from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_word_doc():
    doc = Document()
    
    # Set default style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # Add Title
    title = doc.add_heading('TỔNG HỢP PHÂN TÍCH NHÓM NGƯỜI DÙNG\nỨNG DỤNG HỖ TRỢ NGƯỜI ĂN CHAY', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph("Dưới đây là bản tổng hợp, chuẩn hóa các đặc điểm và kỳ vọng của từng nhóm đối tượng người dùng (User Classes) và hệ thống, dựa trên sự kết hợp các bài phân tích.")
    doc.add_paragraph("")

    # Define the synthesized data
    user_classes = [
        {
            "name": "1. Người ăn chay đã có kinh nghiệm (Vegetarian User)",
            "characteristics": [
                "Là những người đã có thói quen ăn chay, quan tâm đến chế độ ăn thuần thực vật.",
                "Sử dụng hệ thống thường xuyên để tìm kiếm công thức, địa điểm ăn uống, chia sẻ kinh nghiệm và kết nối cộng đồng.",
                "Đa dạng về độ tuổi và trình độ sử dụng công nghệ (từ cơ bản đến thành thạo).",
                "Có nhu cầu cá nhân hóa cao đối với các thông số sức khỏe riêng biệt (chỉ số BMI, mục tiêu sức khỏe, dị ứng thực phẩm)."
            ],
            "expectations": [
                "Tìm kiếm công thức nấu ăn ngon, dễ làm và nhanh chóng.",
                "Được AI gợi ý thực đơn chính xác, phù hợp với sở thích, nguyên liệu sẵn có (theo mùa/vùng miền) và tình trạng sức khỏe.",
                "Tìm kiếm các quán ăn/nhà hàng chay gần vị trí hiện tại một cách dễ dàng và chính xác.",
                "Có thể tìm và mua nguyên liệu chay trực tiếp dựa trên công thức đã chọn.",
                "Theo dõi lượng dinh dưỡng hàng ngày, có thể đồng bộ dữ liệu với các thiết bị đeo/ứng dụng sức khỏe bên thứ ba (Apple Health/Google Fit).",
                "Tham gia vào một cộng đồng sôi nổi; dễ dàng quản lý (đăng, sửa, xóa) bài viết, vlog, bình luận của cá nhân.",
                "Được đảm bảo bảo mật tuyệt đối về các thông tin cá nhân và dữ liệu sức khỏe nhạy cảm."
            ]
        },
        {
            "name": "2. Người mới ăn chay (New Vegetarian / Beginner)",
            "characteristics": [
                "Người mới bắt đầu tìm hiểu hoặc mới chuyển sang chế độ ăn chay, còn thiếu kinh nghiệm.",
                "Chưa biết cách xây dựng một chế độ ăn phù hợp, đảm bảo sức khỏe.",
                "Cần nhiều sự hướng dẫn chi tiết và động lực để duy trì thói quen mới."
            ],
            "expectations": [
                "Có một lộ trình ăn chay rõ ràng, từng bước được cá nhân hóa cho người mới.",
                "Được hướng dẫn chi tiết về cách cân bằng dinh dưỡng, tránh thiếu hụt chất.",
                "Tiếp cận được với cộng đồng thân thiện, sẵn sàng hỗ trợ, giải đáp thắc mắc và truyền kinh nghiệm.",
                "Trải nghiệm giao diện ứng dụng đơn giản, trực quan và dễ sử dụng."
            ]
        },
        {
            "name": "3. Khách vãng lai / Người dùng chưa đăng ký (Guest / Unauthorized User)",
            "characteristics": [
                "Là người dùng đang trong giai đoạn tìm hiểu, dùng thử ứng dụng, chưa xác định rõ nhu cầu.",
                "Tần suất truy cập thấp, không thường xuyên (truy cập ngẫu nhiên hoặc lần đầu).",
                "Trình độ công nghệ đa dạng, mục tiêu chính là đánh giá giá trị của ứng dụng trước khi quyết định tạo tài khoản."
            ],
            "expectations": [
                "Được xem và tìm kiếm các nội dung công khai (video, blog cơ bản) mà không bắt buộc phải đăng nhập.",
                "Giao diện không yêu cầu đăng nhập ngay, nhưng dễ dàng chuyển sang luồng đăng ký khi cần.",
                "Trải nghiệm thử AI Chatbot ở mức độ cơ bản để thấy được tính hữu ích của ứng dụng.",
                "Được thông báo rõ ràng khi chạm đến giới hạn tính năng của người dùng chưa đăng ký (ví dụ: số lượt hỏi AI).",
                "Quy trình đăng ký tài khoản mới phải diễn ra nhanh gọn, đơn giản, không qua quá nhiều bước."
            ]
        },
        {
            "name": "4. Người chia sẻ công thức / KOL (Recipe Contributor / Creator)",
            "characteristics": [
                "Những người có đam mê với ẩm thực chay, thích sáng tạo và nấu nướng.",
                "Thường xuyên đóng góp nội dung, công thức, video/vlog chất lượng cho cộng đồng."
            ],
            "expectations": [
                "Công thức và bài đăng tiếp cận được nhiều người, thu hút lượt xem và tương tác cao.",
                "Cung cấp giao diện đăng bài/chỉnh sửa nội dung trực quan, tiện lợi.",
                "Được cộng đồng và hệ thống ghi nhận những đóng góp (thông qua hệ thống huy hiệu, danh hiệu hoặc verified badge)."
            ]
        },
        {
            "name": "5. Chuyên gia dinh dưỡng (Nutritionist)",
            "characteristics": [
                "Những người có chuyên môn sâu về y tế, dinh dưỡng và chế độ ăn thuần thực vật.",
                "Đóng vai trò cung cấp kiến thức chuẩn xác, tư vấn và kiểm chứng thông tin trên nền tảng."
            ],
            "expectations": [
                "Các bài viết, tài liệu chuyên môn được ưu tiên hiển thị, gắn nhãn đáng tin cậy.",
                "Được cấp tài khoản có xác minh chuyên môn (Verified badge) rõ ràng.",
                "Cung cấp công cụ trả lời câu hỏi, tư vấn dinh dưỡng cho người dùng một cách hiệu quả.",
                "Có quyền hạn tham gia đánh giá, góp ý thực đơn và hỗ trợ hệ thống kiểm duyệt nội dung dinh dưỡng do cộng đồng chia sẻ."
            ]
        },
        {
            "name": "6. Chủ nhà hàng / Quán ăn chay (Restaurant / Food Business Owner)",
            "characteristics": [
                "Các cá nhân, tổ chức đang sở hữu hoặc vận hành các cơ sở kinh doanh dịch vụ ăn uống chay.",
                "Sử dụng nền tảng như một kênh tiếp thị để quảng bá thương hiệu."
            ],
            "expectations": [
                "Đăng tải, cập nhật thông tin quán (menu, địa chỉ, giờ mở cửa) nhanh chóng.",
                "Tiếp cận được tệp khách hàng tiềm năng có nhu cầu ăn chay tại địa phương.",
                "Có công cụ để theo dõi thống kê và phản hồi trực tiếp các bình luận, đánh giá từ thực khách nhằm cải thiện dịch vụ."
            ]
        },
        {
            "name": "7. Quản trị viên (Administrator)",
            "characteristics": [
                "Nhân viên vận hành cốt lõi của hệ thống, tần suất sử dụng hàng ngày.",
                "Có kiến thức kỹ thuật cơ bản (đọc hiểu số liệu, log của AI) nhưng không nhất thiết phải là chuyên gia AI sâu.",
                "Sở hữu quyền truy cập cao nhất, quản lý toàn bộ tài khoản, nội dung và giám sát AI."
            ],
            "expectations": [
                "Sở hữu Dashboard tổng quan rõ ràng, trực quan để theo dõi các số liệu thống kê (bài viết, người dùng, hoạt động AI).",
                "Phát hiện nhanh và có công cụ xử lý các nội dung vi phạm, báo cáo/khiếu nại (đặc biệt là nội dung được AI gắn cờ).",
                "Kiểm soát, theo dõi audit log, đánh giá độ chính xác của các mô hình AI để tinh chỉnh khi cần thiết.",
                "Quản lý linh hoạt danh mục (category) món ăn, công thức.",
                "Đảm bảo và duy trì sự hoạt động ổn định của toàn hệ thống."
            ]
        },
        {
            "name": "8. Hệ thống AI (AI System / Agent)",
            "characteristics": [
                "Là một tác nhân hoạt động tự động và liên tục (24/7).",
                "Bao gồm 5 module chính: Recommendation AI (Gợi ý), AI Nutrition Chatbot (Tư vấn LLM), Computer Vision AI (Nhận diện nguyên liệu), Content Moderation AI (Kiểm duyệt), Speech-to-Text + Summarization AI (Tóm tắt video).",
                "Cần nhận dữ liệu đầu vào (ảnh, câu hỏi, BMI) để xử lý. Có khả năng tự học và cải thiện qua thời gian.",
                "Đóng vai trò là cầu nối giao tiếp giữa dữ liệu người dùng và kho nội dung hệ thống. Độ chính xác không tuyệt đối, vẫn cần sự giám sát của Admin."
            ],
            "expectations": [
                "Xử lý và trả về kết quả (gợi ý, nhận diện, tóm tắt) nhanh chóng, mượt mà, không làm gián đoạn trải nghiệm người dùng.",
                "AI Chatbot phải phản hồi tự nhiên, giải thích các thuật ngữ y khoa/dinh dưỡng bằng ngôn ngữ thông thường dễ hiểu.",
                "Tự động nhận diện nguyên liệu qua ảnh, đánh giá độ tươi và gợi ý công thức tương ứng.",
                "Tự động tóm tắt video/blog dài thành các bước nấu ăn ngắn gọn.",
                "Tự động nhận diện nội dung nhạy cảm/vi phạm và gắn cờ để Admin duyệt thủ công (không tự ý xóa hoàn toàn khi không chắc chắn).",
                "Minh bạch trong hoạt động: độ chính xác phải được đo lường và lưu trữ log cho Administrator kiểm tra."
            ]
        }
    ]

    # Add content
    for uc in user_classes:
        # Heading 2 for Class Name
        h2 = doc.add_heading(uc['name'], level=2)
        h2_font = h2.runs[0].font
        h2_font.color.rgb = RGBColor(46, 116, 181) # A nice blue color

        # Characteristics
        doc.add_heading('Đặc điểm:', level=3)
        for char in uc['characteristics']:
            doc.add_paragraph(char, style='List Bullet')
            
        # Expectations
        doc.add_heading('Kỳ vọng:', level=3)
        for exp in uc['expectations']:
            doc.add_paragraph(exp, style='List Bullet')
            
        doc.add_paragraph("") # Space between sections

    # Save the document
    file_path = '/mnt/data/Tong_Hop_Phan_Tich_Nguoi_Dung_App_An_Chay.docx'
    doc.save(file_path)
    return file_path

create_word_doc()